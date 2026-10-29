from pathlib import Path

import pytest

import informe_docx as gen

SITE_PAGES = {"informe/1-2-1-objetivos-y-alcances.md", "adr/0019-comercio-perfil-no-rol.md"}

NAV = """\
site_name: x
nav:
  - Inicio: index.md
  - 1 Desarrollo:
      - informe/1-1-1-competidores.md
      # comentario: - informe/1-2-3-pantallas.md
      - informe/1-2-4-modelo-de-datos.md
  - Decisiones:
      - adr/README.md
      - adr/0001-monorepo.md
  - Registro: informe/README.md
"""


# repo de prueba
@pytest.fixture
def repo(tmp_path: Path) -> Path:
    (tmp_path / "docs/informe/img").mkdir(parents=True)
    (tmp_path / "docs/informe/img/foto.png").write_bytes(b"png")
    (tmp_path / "docs/adr").mkdir(parents=True)
    (tmp_path / "mkdocs.yml").write_text(NAV, encoding="utf-8")
    return tmp_path


def test_nav_files_keeps_order_and_skips_comments(repo):
    assert gen.nav_files(repo) == [
        "informe/1-1-1-competidores.md",
        "informe/1-2-4-modelo-de-datos.md",
        "adr/README.md",
        "adr/0001-monorepo.md",
        "informe/README.md",
    ]


def test_strip_front_matter_removes_metadata_block():
    md = '---\nseccion: "1.2.4"\n---\n\n# 1.2.4 Modelo de Datos\n'
    assert gen.strip_front_matter(md) == "\n# 1.2.4 Modelo de Datos\n"


def test_root_number_is_first_numbered_heading():
    assert gen.root_number("intro\n\n# 1.2.4 Modelo de Datos\n## 1.2.4.1 X") == "1.2.4"


def test_root_number_none_without_numbered_heading():
    assert gen.root_number("# PATITAS\n") is None


def test_shift_headings_moves_all_levels_by_first_heading_offset():
    md = "# 1.2.4 Modelo\n\n## Entidades\n\n### Detalle"
    assert gen.shift_headings(md, 3) == "### 1.2.4 Modelo\n\n#### Entidades\n\n##### Detalle"


def test_shift_headings_ignores_hash_lines_inside_code_fences():
    md = "# 1.3 X\n\n```bash\n# comentario\n```"
    assert gen.shift_headings(md, 2) == "## 1.3 X\n\n```bash\n# comentario\n```"


def test_splice_replaces_section_until_next_heading_of_same_level():
    md = "# 1.2.2 A\n\n## 1.2.2.2 Historias\n\nVer otro archivo.\n\n## 1.2.2.3 Tareas\n\nT"
    out = gen.splice(md, "1.2.2.2", "SUB")
    assert "Ver otro archivo." not in out
    assert out.index("SUB") < out.index("## 1.2.2.3 Tareas")


def test_splice_leaves_heading_in_its_own_paragraph():
    md = "# 1.2.2 A\n\n## 1.2.2.2 Historias\n\nx\n\n## 1.2.2.3 Tareas"
    out = gen.splice(md, "1.2.2.2", "- ultima viñeta")
    assert "- ultima viñeta\n\n## 1.2.2.3 Tareas" in out


def test_fix_links_page_in_document_becomes_plain_text(repo):
    src = repo / "docs/informe/1-2-2-modelo-funcional.md"
    md = "Ver [1.2.1](1-2-1-objetivos-y-alcances.md#objetivos)."
    assert gen.fix_links(md, src, repo, SITE_PAGES) == "Ver 1.2.1."


def _write(repo: Path, rel: str, text: str) -> Path:
    path = repo / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def test_expand_refs_inlines_cited_md_at_end_of_its_section(repo):
    _write(repo, "docs/QA.md", "# QA pipeline\n\nContenido QA.\n\n## Casos\n\nC1")
    src = repo / "docs/informe/1-3-1-x.md"
    md = "# 1.3.1 X\n\n## Calidad\n\nVer [QA](../QA.md) y sigue.\n\nOtro párrafo.\n\n## Otra\n\nfin"
    out = gen.expand_refs(md, src, repo, SITE_PAGES, set())
    assert "Ver QA y sigue." in out
    assert out.index("Otro párrafo.") < out.index("### QA pipeline") < out.index("## Otra")
    assert "#### Casos" in out and "Contenido QA." in out


def test_expand_refs_inlines_each_file_once(repo):
    _write(repo, "docs/QA.md", "# QA\n\nUna vez.")
    src = repo / "docs/informe/x.md"
    md = "# 1 X\n\n[a](../QA.md)\n\n## Y\n\n[b](../QA.md)"
    out = gen.expand_refs(md, src, repo, SITE_PAGES, set())
    assert out.count("Una vez.") == 1


def test_expand_refs_follows_nested_refs_without_cycles(repo):
    _write(repo, "docs/A.md", "# A\n\nA cita [B](B.md).")
    _write(repo, "docs/B.md", "# B\n\nB cita [A](A.md).\n\n![F](img/f.png)")
    (repo / "docs/img").mkdir()
    (repo / "docs/img/f.png").write_bytes(b"png")
    src = repo / "docs/informe/x.md"
    out = gen.expand_refs("# 1 X\n\n[A](../A.md)", src, repo, SITE_PAGES, set())
    assert out.count("A cita B.") == 1 and out.count("B cita A.") == 1
    assert "### B" in out
    assert f"![F]({(repo / 'docs/img/f.png').as_posix()})" in out


def test_expand_refs_does_not_inline_excluded_files(repo):
    _write(repo, "docs/adr/0000-plantilla.md", "# Plantilla\n\nNo va.")
    src = repo / "docs/adr/README.md"
    out = gen.expand_refs("# ADR\n\n[plantilla](0000-plantilla.md)", src, repo, SITE_PAGES, set())
    assert "No va." not in out and "plantilla" in out


def test_shift_headings_turns_levels_beyond_six_into_bold():
    assert gen.shift_headings("# A\n\n## B", 6) == "###### A\n\n**B**"


def test_fix_links_other_repo_file_goes_to_github_blob(repo):
    src = repo / "docs/informe/1-2-2-modelo-funcional.md"
    md = "[PDD](../pdd/PDD%20v1.0.docx)"
    assert gen.fix_links(md, src, repo, SITE_PAGES) == (
        "[PDD](https://github.com/ISI-ProyectoFinal-2026/Grupo1/blob/main/docs/pdd/PDD%20v1.0.docx)"
    )


def test_fix_links_local_image_becomes_absolute_path(repo):
    src = repo / "docs/informe/1-1-1-competidores.md"
    out = gen.fix_links("![Foto](img/foto.png)", src, repo, SITE_PAGES)
    assert out == f"![Foto]({(repo / 'docs/informe/img/foto.png').as_posix()})"


def test_fix_links_keeps_external_and_anchor_links(repo):
    src = repo / "docs/informe/1-1-1-competidores.md"
    md = "[a](https://example.com) [b](#seccion)"
    assert gen.fix_links(md, src, repo, SITE_PAGES) == md


def test_report_files_excludes_registry_and_template(repo):
    files = gen.nav_files(repo) + ["adr/0000-plantilla.md"]
    informe, adrs = gen.report_files(files)
    assert informe == ["informe/1-1-1-competidores.md", "informe/1-2-4-modelo-de-datos.md"]
    assert adrs == ["adr/0001-monorepo.md"]


def test_default_output_name_has_date():
    assert gen.output_name("2026-10-28") == "Patitas-Informe-Final-2026-10-28.docx"


TABLE_MD = """\
Texto previo.

| A | B |
|---|---|
| 1 | 2 |

Tabla 1.2.1-1. Objetivos del PDD contrastados con el código. Fuente: PDD v1.0 y `git ls-files`.

Texto posterior."""


def test_caption_of_table_moves_above_with_bold_number_and_italic_title():
    out = gen.format_captions(TABLE_MD)
    caption = '::: {custom-style="Epígrafe tabla"}\n**Tabla 1.2.1-1**\\\n*Objetivos del PDD contrastados con el código*\n:::'
    assert caption in out
    assert out.index(caption) < out.index("| A | B |")


def test_caption_source_becomes_note_below_table():
    out = gen.format_captions(TABLE_MD)
    note = '::: {custom-style="Nota"}\n*Nota.* Fuente: PDD v1.0 y `git ls-files`.\n:::'
    assert note in out
    assert out.index("| 1 | 2 |") < out.index(note) < out.index("Texto posterior.")
    assert "Tabla 1.2.1-1. Objetivos" not in out


def test_caption_of_image_moves_above_image():
    md = "![Alt](img/a.png)\n\nFigura 1.1-1. Perfil de Instagram. Fuente: `img/a.png`.\n"
    out = gen.format_captions(md)
    assert out.index('custom-style="Epígrafe figura"') < out.index("![Alt](img/a.png)")
    assert out.index("![Alt](img/a.png)") < out.index("*Nota.* Fuente: `img/a.png`.")


def test_caption_without_table_or_image_is_left_untouched():
    md = "Párrafo.\n\nTabla 9-1. Huérfana. Fuente: nada.\n"
    assert gen.format_captions(md) == md


def test_caption_inside_code_fence_is_ignored():
    md = "```\n| a |\n\nTabla 1-1. X. Fuente: y.\n```"
    assert gen.format_captions(md) == md


def test_cover_lines_include_faculty_career_and_teachers():
    lines = gen.cover_lines("2026-10-28")
    assert "Universidad Tecnológica Nacional" in lines
    assert "Facultad Regional San Rafael" in lines
    assert "Ingeniería en Sistemas de Información" in lines
    assert "Docentes: Hernán Zinik y Matías Varela" in lines
    assert "28 de octubre de 2026" in lines
    assert not any("[COMPLETAR" in line for line in lines)


def _doc_with_table():
    import docx

    d = docx.Document()
    d.add_paragraph("Uno", style="Heading 1")
    t = d.add_table(rows=3, cols=2)
    for row in t.rows:
        for cell in row.cells:
            cell.text = "x"
    return d


def test_apa_table_has_only_top_header_and_bottom_rules():
    d = _doc_with_table()
    gen.format_tables(d)
    xml = d.tables[0]._tbl.xml
    for side in ("top", "bottom"):
        assert f'<w:{side} w:val="single"' in xml
    for side in ("insideV", "insideH", "left", "right"):
        assert f'<w:{side} w:val="nil"' in xml
    assert "<w:tblHeader/>" in d.tables[0].rows[0]._tr.xml


def test_page_number_field_goes_in_header_right():
    d = _doc_with_table()
    gen.add_page_number(d)
    header = d.sections[0].header
    assert "PAGE" in header._element.xml
    assert header.paragraphs[0].alignment == 2


def test_heading_1_starts_on_new_page():
    d = _doc_with_table()
    gen.format_styles(d)
    assert d.styles["Heading 1"].paragraph_format.page_break_before is True
