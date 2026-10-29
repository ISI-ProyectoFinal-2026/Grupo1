# genera el informe final en .docx con el esqueleto y los estilos del template
import argparse
import datetime
import hashlib
import re
import shutil
import subprocess
from pathlib import Path
from urllib.parse import quote, unquote

REPO_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = "docs/Template Presentacion Final (1).docx"
BLOB_URL = "https://github.com/ISI-ProyectoFinal-2026/Grupo1/blob/main/"
EXCLUDED = {"informe/README.md", "adr/README.md", "adr/0000-plantilla.md"}
DECISIONS = "2.2.2.4"

# integrantes segun docs/pdd/PDD v1.0.docx
TEAM = [
    "Alvarado, Matías",
    "Diaz Bellomo, Emiliano Agustín",
    "Marín, Franco",
    "Valdivieso, Gabriel",
]
MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
          "septiembre", "octubre", "noviembre", "diciembre"]
CAPTION_STYLES = {"Tabla": "Epígrafe tabla", "Figura": "Epígrafe figura"}
NOTE_STYLE = "Nota"
LISTS = {"5 Índice de figuras": "Epígrafe figura", "6 Índice de tablas": "Epígrafe tabla"}
TOC_DEPTH = 5

NUM = re.compile(r"^(\d+(?:\.\d+)*)\s")
CAPTION = re.compile(r"^(Tabla|Figura) (\d[\d.]*-\d+)\. (.+?)\. Fuente: (.+)$")
HEAD = re.compile(r"^(#{1,9})\s+(.*)$")
MD_LINK = re.compile(r"(!?)\[([^\]]*)\]\(([^)\s]+)\)")
MAX_HEADING = 6
MERMAID = re.compile(r"^```mermaid\s*\n(.*?)^```\s*$", re.S | re.M)
NAV_ENTRY = re.compile(r"^\s*-\s*(?:[^:\n#]+:\s*)?((?:informe|adr)/[\w.-]+\.md)\s*$", re.M)


def output_name(date: str) -> str:
    return f"Patitas-Informe-Final-{date}.docx"


def nav_files(repo: Path) -> list[str]:
    text = (repo / "mkdocs.yml").read_text(encoding="utf-8")
    return NAV_ENTRY.findall(text[text.index("\nnav:"):])


def report_files(files: list[str]) -> tuple[list[str], list[str]]:
    kept = [f for f in files if f not in EXCLUDED]
    return [f for f in kept if f.startswith("informe/")], [f for f in kept if f.startswith("adr/")]


def strip_front_matter(md: str) -> str:
    return re.sub(r"\A---\n.*?\n---\n", "", md, flags=re.S)


def root_number(md: str) -> str | None:
    for line in md.splitlines():
        h = HEAD.match(line)
        if h:
            n = NUM.match(h.group(2))
            return n.group(1) if n else None
    return None


def level_of(number: str) -> int:
    return number.count(".") + 1


def shift_headings(md: str, level_of_first: int) -> str:
    out, offset, fence = [], None, False
    for line in md.splitlines():
        if line.startswith("```"):
            fence = not fence
        h = None if fence else HEAD.match(line)
        if h:
            if offset is None:
                offset = level_of_first - len(h.group(1))
            level = len(h.group(1)) + offset
            line = "#" * level + " " + h.group(2) if level <= MAX_HEADING else f"**{h.group(2)}**"
        out.append(line)
    return "\n".join(out)


def splice(md: str, number: str, sub_md: str) -> str:
    # reemplaza la seccion `number` (titulo y cuerpo) por otro contenido
    lines = md.splitlines()
    start = end = depth = None
    for i, line in enumerate(lines):
        h = HEAD.match(line)
        if not h:
            continue
        if start is None and h.group(2).startswith(number + " "):
            start, depth = i, len(h.group(1))
        elif start is not None and len(h.group(1)) <= depth:
            end = i
            break
    if start is None:
        return md
    end = len(lines) if end is None else end
    return "\n".join(lines[:start] + ["", sub_md, ""] + lines[end:])


def fix_links(md: str, src: Path, repo: Path, site_pages: set[str], cited: list[Path] | None = None) -> str:
    # documento autocontenido: los .md citados quedan como texto y se desglosan (expand_refs)
    def repl(m: re.Match) -> str:
        bang, text, raw = m.groups()
        if re.match(r"^[a-z]+:|^#", raw):
            return m.group(0)
        path, _, frag = raw.partition("#")
        target = (src.parent / unquote(path)).resolve()
        if bang and target.suffix.lower() == ".png" and target.exists():
            return f"![{text}]({target.as_posix()})"
        rel = target.relative_to(repo.resolve()).as_posix()
        if not bang and target.suffix.lower() == ".md":
            docs_rel = rel[5:] if rel.startswith("docs/") else None
            if cited is not None and target.exists() and docs_rel not in site_pages and docs_rel not in EXCLUDED:
                cited.append(target)
            return text
        anchor = "#" + frag if frag else ""
        return f"{bang}[{text}]({BLOB_URL}{quote(rel)}{anchor})"

    return MD_LINK.sub(repl, md)


def expand_refs(md: str, src: Path, repo: Path, site_pages: set[str], seen: set[Path]) -> str:
    # cada .md citado fuera del documento va al final de la seccion que lo cita, una sola vez
    out, pending, level, fence = [], [], 1, False

    def flush() -> None:
        for target in pending:
            sub = strip_front_matter(target.read_text(encoding="utf-8"))
            out.extend(["", shift_headings(expand_refs(sub, target, repo, site_pages, seen), level + 1), ""])
        pending.clear()

    for line in md.splitlines():
        if line.startswith("```"):
            fence = not fence
        h = None if fence else HEAD.match(line)
        if h:
            flush()
            level = len(h.group(1))
        if not fence:
            cited: list[Path] = []
            line = fix_links(line, src, repo, site_pages, cited)
            for target in cited:
                if target not in seen:
                    seen.add(target)
                    pending.append(target)
        out.append(line)
    flush()
    return "\n".join(out)


def render_mermaid(md: str, work: Path, mmdc: str) -> str:
    def repl(m: re.Match) -> str:
        code = m.group(1)
        name = hashlib.sha1(code.encode()).hexdigest()[:12]
        src, png = work / f"{name}.mmd", work / f"{name}.png"
        if not png.exists():
            src.write_text(code, encoding="utf-8")
            subprocess.run([mmdc, "-i", str(src), "-o", str(png), "-s", "2", "-b", "white"],
                           check=True, capture_output=True)
        return f"![]({png.as_posix()})\n"

    return MERMAID.sub(repl, md)


def format_captions(md: str) -> str:
    # epigrafe APA 7: numero y titulo arriba, fuente como nota debajo
    lines = md.split("\n")
    fenced, inside = [], False
    for line in lines:
        if line.startswith("```"):
            inside = not inside
            fenced.append(True)
        else:
            fenced.append(inside)
    out = list(lines)
    for i in range(len(lines) - 1, -1, -1):
        m = None if fenced[i] else CAPTION.match(lines[i])
        if not m:
            continue
        j = i - 1
        while j >= 0 and not lines[j].strip():
            j -= 1
        if j < 0 or fenced[j]:
            continue
        if lines[j].startswith("|"):
            start = j
            while start > 0 and lines[start - 1].startswith("|"):
                start -= 1
        elif lines[j].startswith("!["):
            start = j
        else:
            continue
        kind, number, title, source = m.groups()
        out[i:i + 1] = [f'::: {{custom-style="{NOTE_STYLE}"}}', f"*Nota.* Fuente: {source}", ":::"]
        out[start:start] = ["", f'::: {{custom-style="{CAPTION_STYLES[kind]}"}}',
                            f"**{kind} {number}**\\", f"*{title}*", ":::", ""]
    return "\n".join(out)


def cover_lines(date: str) -> list[str]:
    year, month, day = (int(x) for x in date.split("-"))
    return [
        "Universidad Tecnológica Nacional",
        "Facultad Regional San Rafael",
        "Ingeniería en Sistemas de Información",
        "Proyecto Final 2026",
        "PATITAS",
        "Informe Final",
        "Grupo 1",
        *TEAM,
        "Docentes: Hernán Zinik y Matías Varela",
        f"{day} de {MONTHS[month - 1]} de {year}",
    ]


def _w(tag: str, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement(f"w:{tag}")
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), str(value))
    return el


def _insert_ordered(parent, child, successors: tuple[str, ...]):
    # respeta el orden del esquema OOXML
    from docx.oxml.ns import qn

    for old in parent.findall(child.tag):
        parent.remove(old)
    for el in parent:
        if el.tag in {qn(f"w:{s}") for s in successors}:
            el.addprevious(child)
            return child
    parent.append(child)
    return child


def _field(paragraph, instruction: str, placeholder: str = ""):
    for kind in ("begin", None, "separate", placeholder, "end"):
        run = paragraph.add_run()
        if kind in ("begin", "separate", "end"):
            run._r.append(_w("fldChar", fldCharType=kind))
        elif kind is None:
            instr = _w("instrText")
            instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            instr.text = f" {instruction} "
            run._r.append(instr)
        elif kind:
            run.text = kind


def _style(doc, name: str):
    from docx.enum.style import WD_STYLE_TYPE

    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        return style


def format_styles(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    for name in CAPTION_STYLES.values():
        pf = _style(doc, name).paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before, pf.space_after, pf.line_spacing = Pt(12), Pt(6), 1.0
        pf.keep_with_next = True
        _style(doc, name).font.color.rgb = RGBColor(0, 0, 0)
    note = _style(doc, NOTE_STYLE)
    note.font.size = Pt(10)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_before, note.paragraph_format.space_after = Pt(4), Pt(12)
    note.paragraph_format.line_spacing = 1.0


def format_tables(doc) -> None:
    from docx.shared import Pt

    border = {"val": "single", "sz": 8, "space": 0, "color": "000000"}
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = _w("tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_w(side, **border) if side in ("top", "bottom") else _w(side, val="nil"))
        _insert_ordered(tbl_pr, borders, ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"))
        _insert_ordered(tbl_pr, _w("jc", val="left"), ("tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
                                                         "tblCellMar", "tblLook", "tblCaption", "tblDescription"))
        for index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            _insert_ordered(tr_pr, _w("cantSplit"), ("trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"))
            if index == 0:
                _insert_ordered(tr_pr, _w("tblHeader"), ("tblCellSpacing", "jc", "hidden"))
            for cell in row.cells:
                if index == 0:
                    tc_borders = _w("tcBorders")
                    tc_borders.append(_w("bottom", val="single", sz=4, space=0, color="000000"))
                    _insert_ordered(cell._tc.get_or_add_tcPr(), tc_borders,
                                    ("shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"))
                for paragraph in cell.paragraphs:
                    pf = paragraph.paragraph_format
                    pf.space_before, pf.space_after, pf.line_spacing = Pt(2), Pt(2), 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if index == 0:
                            run.font.bold = True


def format_images(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True


def add_cover(doc, date: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    lines = cover_lines(date)
    sizes = {"PATITAS": None, "Informe Final": Pt(20)}
    body = doc.element.body
    anchor = body[0]
    for index, text in enumerate(lines):
        style = "Title" if text == "PATITAS" else None
        paragraph = doc.add_paragraph(text, style=style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.space_after = Pt(6)
        if index == 0:
            pf.space_before = Pt(60)
        if text in ("PATITAS", "Grupo 1"):
            pf.space_before = Pt(60)
        if text.startswith("Docentes"):
            pf.space_before = Pt(36)
        for run in paragraph.runs:
            run.font.size = sizes.get(text, Pt(12))
        if index == len(lines) - 1:
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        anchor.addprevious(paragraph._p)


def add_lists(doc) -> None:
    for paragraph in list(doc.paragraphs):
        style = LISTS.get(paragraph.text.strip())
        if style and paragraph.style.name == "Heading 1":
            target = doc.add_paragraph()
            _field(target, f'TOC \\h \\z \\t "{style},1"')
            paragraph._p.addnext(target._p)


def add_page_number(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _field(paragraph, "PAGE", "1")


def postprocess(out: Path, date: str) -> None:
    import docx

    doc = docx.Document(str(out))
    format_styles(doc)
    format_tables(doc)
    format_images(doc)
    add_lists(doc)
    add_cover(doc, date)
    add_page_number(doc)
    doc.save(str(out))


WORD_UPDATE = r"""
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
$d = $w.Documents.Open($args[0], $false, $false, $false)
$sep = (Get-Culture).TextInfo.ListSeparator
foreach ($f in $d.Fields) {
    if ($f.Code.Text -match '\\t "') { $f.Code.Text = $f.Code.Text -replace ',1"', ($sep + '1"') }
}
foreach ($t in $d.TablesOfContents) { $t.Update() | Out-Null }
$d.Fields.Update() | Out-Null
foreach ($t in $d.TablesOfContents) { $t.Update() | Out-Null }
$d.Save()
$d.Close()
$w.Quit()
"""


def update_fields(out: Path) -> bool:
    # Word completa indices y numeros de pagina; Drive no actualiza campos
    import docx

    script = out.parent / ".informe-docx-work" / "update.ps1"
    script.write_text(WORD_UPDATE, encoding="utf-8")
    result = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), str(out)],
                            capture_output=True, text=True)
    if result.returncode != 0:
        print("Word no disponible: los indices se completan al abrir el archivo en Word.", result.stderr.strip())
        return False
    doc = docx.Document(str(out))
    for el in doc.settings.element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields"):
        doc.settings.element.remove(el)
    doc.save(str(out))
    return True


def template_skeleton(template: Path) -> list[tuple[int, str]]:
    import docx

    items = []
    for p in docx.Document(str(template)).paragraphs:
        name, text = p.style.name, p.text.strip()
        if name.startswith("Heading ") and text:
            items.append((int(name.split()[1]), text))
    return items


def assemble(repo: Path, work: Path, mmdc: str) -> str:
    files = nav_files(repo)
    site_pages = set(files)
    informe, adrs = report_files(files)
    seen = {(repo / "docs" / f).resolve() for f in files}

    def load(rel: str) -> str:
        src = repo / "docs" / rel
        md = strip_front_matter(src.read_text(encoding="utf-8"))
        return format_captions(render_mermaid(expand_refs(md, src, repo, site_pages, seen), work, mmdc))

    sections = {}
    for rel in informe:
        md = load(rel)
        sections[root_number(md)] = md
    nested = {o for n in sections for o in sections if o != n and o.startswith(n + ".")}
    for num in sections:
        for other in nested:
            if other.startswith(num + "."):
                sections[num] = splice(sections[num], other, f"<<{other}>>")

    def emit(num: str) -> str:
        md = shift_headings(sections[num], level_of(num))
        for other in nested:
            if f"<<{other}>>" in md:
                md = md.replace(f"<<{other}>>", emit(other))
        return md

    decisions = [shift_headings(load("adr/README.md"), level_of(DECISIONS) + 1)]
    decisions += [shift_headings(load(rel), level_of(DECISIONS) + 1) for rel in adrs]

    parts, owned = [], None
    for level, text in template_skeleton(repo / TEMPLATE):
        n = NUM.match(text)
        num = n.group(1) if n else None
        if owned and num and (num == owned or num.startswith(owned + ".")):
            continue
        owned = None
        if num in sections and num not in nested:
            parts.append(emit(num))
            owned = num
            continue
        parts.append("#" * level + " " + text)
        if num == DECISIONS:
            parts.extend(decisions)
    return "\n\n".join(parts)


def build(repo: Path, out: Path, mmdc: str, date: str, word: bool = True) -> Path:
    import pypandoc

    work = out.parent / ".informe-docx-work"
    work.mkdir(parents=True, exist_ok=True)
    body = assemble(repo, work, mmdc)
    (work / "informe.md").write_text(body, encoding="utf-8")
    pypandoc.convert_text(
        body, "docx", format="markdown+pipe_tables+grid_tables-auto_identifiers-implicit_figures",
        outputfile=str(out),
        extra_args=["--reference-doc", str(repo / TEMPLATE), "--toc", f"--toc-depth={TOC_DEPTH}",
                    "--resource-path", str(repo / "docs"),
                    "-M", "toc-title=Índice de contenidos", "-M", "lang=es-AR"],
    )
    postprocess(out, date)
    if word:
        update_fields(out)
    return out


def main() -> None:
    today = datetime.date.today().isoformat()
    parser = argparse.ArgumentParser(description="Genera el informe final en .docx")
    parser.add_argument("--out", type=Path, default=Path.cwd() / output_name(today))
    parser.add_argument("--mmdc", default=shutil.which("mmdc") or "mmdc", help="ejecutable de mermaid-cli")
    parser.add_argument("--no-word", action="store_true", help="no actualizar indices con Word")
    args = parser.parse_args()
    print(build(REPO_ROOT, args.out.resolve(), args.mmdc, today, word=not args.no_word))


if __name__ == "__main__":
    main()
